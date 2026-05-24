import json
import ast
import re
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement






def set_cell_background(cell, fill_hex="2E7D32"):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_valign(cell, align="center"):
    tcPr = cell._element.get_or_add_tcPr()
    valign = OxmlElement('w:vAlign')
    valign.set(qn('w:val'), align)
    tcPr.append(valign)

def force_run_color(run, hex_color):
    """Жесткое внедрение цвета через XML с удалением перекрывающих тем Word"""
    rPr = run._element.get_or_add_rPr()
    color = rPr.find(qn('w:color'))
    if color is None:
        color = OxmlElement('w:color')
        rPr.append(color)
    color.set(qn('w:val'), hex_color)
    for attr in [qn('w:themeColor'), qn('w:themeTint'), qn('w:themeShade')]:
        if attr in color.attrib:
            del color.attrib[attr]

def force_run_font(run, font_name):
    if font_name and str(font_name).strip():
        fname = str(font_name).strip()
        run.font.name = fname
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn('w:ascii'), fname)
        rFonts.set(qn('w:hAnsi'), fname)
        rFonts.set(qn('w:cs'), fname)

def make_collapsible_heading(para, level):
    """Добавляет кликабельную стрелочку сворачивания/разворачивания к заголовку"""
    pPr = para._element.get_or_add_pPr()
    outlineLvl = pPr.find(qn('w:outlineLvl'))
    if outlineLvl is None:
        outlineLvl = OxmlElement('w:outlineLvl')
        pPr.append(outlineLvl)
    outlineLvl.set(qn('w:val'), str(level))

def format_document(input_path, output_path, font_name, font_size, placeholders_input, is_internal_log=True, **kwargs):
    if os.path.exists(output_path):
        try: os.rename(output_path, output_path)
        except OSError: raise PermissionError(f"Файл {output_path} открыт! Закройте его.")

    doc = Document(input_path)
    
    # парсер
    placeholders = {}
    if placeholders_input:
        if isinstance(placeholders_input, dict):
            placeholders = placeholders_input
        elif isinstance(placeholders_input, str):
            txt = placeholders_input.strip()
            if txt and txt not in ("{}", "None"):
                try: placeholders = json.loads(txt)
                except Exception: placeholders = ast.literal_eval(txt)
    
    clean_placeholders = {}
    if isinstance(placeholders, dict):
        for k, v in placeholders.items():
            key = str(k)
            if not key.startswith("{{"): key = f"{{{{{key}}}}}"
            clean_placeholders[key] = str(v)

    # цвета
    corp_green_hex = "2E7D32"
    light_green_hex = "4CAF50" 
    white_hex = "FFFFFF"
    
    for para in doc.paragraphs:
        if clean_placeholders:
            for run in para.runs:
                text = run.text
                text = text.replace('\t', ' ')
                for key, value in clean_placeholders.items():
                    if key in text: text = text.replace(key, value)
                run.text = text

        full_text = para.text.strip()
        if not full_text:
            p = para._element
            p.getparent().remove(p)
            continue
            
        text_upper = full_text.upper()
        
        is_h1 = False
        is_h2 = False
        
        if any(word in text_upper for word in ['ВВЕДЕНИЕ', 'ЗАКЛЮЧЕНИЕ', 'ОСНОВНАЯ ЧАСТЬ', 'РЕЗУЛЬТАТЫ РАБОТЫ']):
            is_h1 = True
        elif any(word in text_upper for word in ['ОПИСАНИЕ ПРОЕКТА', 'КЛЮЧЕВЫЕ ПОКАЗАТЕЛИ', 'ДОСТИГНУТЫЕ РЕЗУЛЬТАТЫ', 'ПЛАН НА СЛЕДУЮЩИЙ']):
            is_h2 = True
        elif re.match(r'^\s*\d+\.\d+', full_text):
            is_h2 = True
        elif re.match(r'^\s*\d+\.', full_text):
            is_h1 = True
        
        if not (is_h1 or is_h2) and para.style and para.style.name:
            sname = para.style.name.lower()
            if 'heading 1' in sname or 'заголовок 1' in sname: is_h1 = True
            elif 'heading 2' in sname or 'заголовок 2' in sname: is_h2 = True

        if is_h1:
            try: para.style = 'Heading 1'
            except: 
                try: para.style = 'Заголовок 1'
                except: pass
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            make_collapsible_heading(para, 0)
        elif is_h2:
            try: para.style = 'Heading 2'
            except: 
                try: para.style = 'Заголовок 2'
                except: pass
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            make_collapsible_heading(para, 1)

        for run in para.runs:
            force_run_font(run, font_name)
            
            if is_h1:
                run.bold = True
                force_run_color(run, corp_green_hex)
                run.font.size = Pt(16)
            elif is_h2:
                run.bold = True
                force_run_color(run, light_green_hex)
                run.font.size = Pt(14)
            else:
                if font_size:
                    try: run.font.size = Pt(int(float(font_size)))
                    except: pass

    for table in doc.tables:
        table.style = 'Table Grid'
        for row_idx, row in enumerate(table.rows):
            is_header_row = (row_idx == 0)
            for cell in row.cells:
                set_cell_valign(cell, "center")
                if is_header_row:
                    set_cell_background(cell, corp_green_hex)
                
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        text = run.text
                        text = text.replace('\t', ' ')
                        if clean_placeholders:
                            for key, value in clean_placeholders.items():
                                if key in text: text = text.replace(key, value)
                        run.text = text
                        
                        force_run_font(run, font_name)
                        if is_header_row:
                            run.bold = True
                            force_run_color(run, white_hex)
                        else:
                            if font_size:
                                try: run.font.size = Pt(int(float(font_size)))
                                except: pass

    doc.save(output_path)
    if is_internal_log: print(f"Успешно сохранено: {output_path}")

format_document(
    ${vc('SOURCE_PATH')}, 
    ${vc('OUTPUT_PATH')}, 
    ${vc('FONT_NAME')}, 
    ${vc('FONT_SIZE')}, 
    ${vc('PLACEHOLDERS')}, 
    is_internal_log=True
)